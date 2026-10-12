#!/usr/bin/env python3
"""Narrative slide + dashboard screenshots → board-ready PPTX."""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from project_paths import DECK_EXPORT_DIR, ROOT, SCREENSHOT_DIR

EAC_TEMPLATE = Path(
    "/Users/aishwaryajagadish/Downloads/2026 EAC.Overall andAdmin Core.9.22.26.pptx"
)
OUTPUT_PPTX = DECK_EXPORT_DIR / "USC_ADRC_Dashboard_Narrative_and_Screenshots.pptx"
OUTPUT_DOWNLOADS = Path.home() / "Desktop" / "USC_ADRC_Dashboard_Narrative_and_Screenshots.pptx"
OUTPUT_SCREENSHOTS_ONLY = DECK_EXPORT_DIR / "USC_ADRC_Dashboard_Screenshots.pptx"
OUTPUT_SCREENSHOTS_DESKTOP = Path.home() / "Desktop" / "USC_ADRC_Dashboard_Screenshots.pptx"
CAPTURE_SCRIPT = ROOT / "scripts" / "capture_dashboard_screenshots.mjs"

SCREENSHOT_SLIDES = [
    ("01-header-recruitment.png", "Recruitment · GeneScreen, ADRC, BRS"),
    ("02-demographics.png", "Demographics · Table 1 charts and cohort table"),
    ("03-neuropathology-specimen.png", "Neuropathology · specimen sharing (neurobank)"),
    ("05-biomarkers-specimen.png", "Biomarkers · biofluid specimen sharing"),
    ("04-ore.png", "ORE outreach"),
    ("06-imaging.png", "Imaging · SCAN & CLARiTI YTD"),
    ("00-full-page.png", "Full dashboard scroll (reference)"),
]

def load_screenshot_manifest() -> list[dict]:
    manifest_path = SCREENSHOT_DIR / "manifest.json"
    if manifest_path.is_file():
        data = json.loads(manifest_path.read_text(encoding="utf-8"))
        return list(data.get("slides", []))
    return [
        {"file": f, "caption": c}
        for f, c in SCREENSHOT_SLIDES
        if not f.startswith("00-full")
    ]

EAC_MARGIN = Inches(0.14)
EAC_CAPTION_BAND = Inches(0.42)
# Metric slides use blank layout (full slide); title slide keeps EAC master footer.
EAC_BOTTOM = Inches(0.12)
NAVY = RGBColor(26, 39, 68)
PURPLE = RGBColor(91, 75, 138)
GRAY = RGBColor(80, 80, 80)


def load_metrics() -> dict:
    pkg = json.loads((ROOT / "src/data/rpprDataPackage.json").read_text(encoding="utf-8"))
    ytd = pkg["reportingPeriodSummary"]
    uds = ytd["udsClinical"]
    redcap = json.loads((ROOT / "src/data/redcapSpecimenSharing.json").read_text(encoding="utf-8"))
    gs = json.loads((ROOT / "src/data/geneScreenYtd.json").read_text(encoding="utf-8"))
    snap = gs["registrySnapshot"]
    bf = redcap["biofluids"]["reportingPeriod"]
    nb = redcap["neurobank"]["reportingPeriod"]
    return {
        "window": ytd["label"],
        "udsParticipants": uds["uniqueParticipantsWithVisit"],
        "initialEvals": uds["newInitialEvaluations"],
        "geneScreenTotal": snap["membersComplete"],
        "geneScreenYtdNew": gs["newMembersByConsentDate"],
        "geneScreenOnline": gs["onlineEnrollmentsByMembershipTimestamp"],
        "bioRequests": bf["requestsReceived"],
        "neuroRequests": nb["requestsReceived"],
        "plasmaAliquots": bf["plasma"]["aliquotsRequested"],
    }


DECK_BUILD_ID = "slide-center-v20"


def restart_demo_for_capture() -> None:
    """Production build + fresh :3002 so Playwright sees latest CSS and DOM."""
    import time
    import urllib.request

    if not os.environ.get("SKIP_DEMO_REBUILD"):
        print("Building dashboard (npm run build)…")
        subprocess.run(["npm", "run", "build"], cwd=str(ROOT), check=True)

    subprocess.run(
        "lsof -ti :3002 | xargs kill -9 2>/dev/null || true",
        shell=True,
        cwd=str(ROOT),
    )
    time.sleep(1)
    subprocess.Popen(
        ["npm", "run", "start:demo"],
        cwd=str(ROOT),
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    for _ in range(80):
        try:
            html = urllib.request.urlopen("http://localhost:3002", timeout=3).read().decode(
                "utf-8", errors="replace"
            )
            if "deck-section-recruitment" in html:
                return
        except OSError:
            time.sleep(0.4)
    raise SystemExit("Dashboard not reachable on :3002 after rebuild. Run: npm run demo")


def capture_screenshots() -> None:
    restart_demo_for_capture()
    if not (ROOT / "node_modules" / "playwright").is_dir():
        raise SystemExit("Run: npm install && npx playwright install chromium")
    result = subprocess.run(
        ["node", str(CAPTURE_SCRIPT), str(SCREENSHOT_DIR)],
        cwd=str(ROOT),
        capture_output=True,
        text=True,
    )
    if result.stdout:
        print(result.stdout.strip())
    if result.returncode != 0:
        raise SystemExit(result.stderr or result.stdout or "Screenshot capture failed")


def add_narrative_slide(prs: Presentation, m: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    w, h = prs.slide_width, prs.slide_height

    title = slide.shapes.add_textbox(Inches(0.6), Inches(0.45), w - Inches(1.2), Inches(0.8))
    tf = title.text_frame
    p = tf.paragraphs[0]
    p.text = "USC ADRC Clinical Core · Aim 1 Metrics Dashboard"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = NAVY

    sub = slide.shapes.add_textbox(Inches(0.6), Inches(1.15), w - Inches(1.2), Inches(0.4))
    sp = sub.text_frame.paragraphs[0]
    sp.text = f"Reporting window: {m['window']}  ·  Mid-cycle RPPR-aligned view"
    sp.font.size = Pt(14)
    sp.font.color.rgb = PURPLE

    body = slide.shapes.add_textbox(Inches(0.6), Inches(1.65), w - Inches(1.2), h - Inches(2.2))
    tf = body.text_frame
    tf.word_wrap = True

    bullets = [
        (
            "Purpose. "
            "Single dashboard for programmatic and core leads—same visuals we use in RPPR, "
            "updated between formal milestone cycles so the center can communicate progress with less back-and-forth."
        ),
        (
            "How to read colors. "
            "Black numbers = calendar year-to-date from the harmonized Data Core package and USC REDCap re-aggregates. "
            "Green = last RPPR / workbook cumulative context. Red = not yet in export (do not cite)."
        ),
        (
            "Highlights this period. "
            f"ADRC {m['udsParticipants']} participants with ≥1 UDS visit; {m['initialEvals']} new initial evaluations. "
            f"GeneScreen registry {m['geneScreenTotal']} members (REDCap); {m['geneScreenYtdNew']} new and "
            f"{m['geneScreenOnline']} online enrollments YTD. "
            f"Specimen sharing: {m['bioRequests']} biofluid requests, {m['plasmaAliquots']} plasma aliquots; "
            f"{m['neuroRequests']} neurobank requests."
        ),
        (
            "Verification. "
            "Automated validation ties on-screen metrics to source JSON/CSV (npm run validate:all). "
            "Attach USC_ADRC_Dashboard_Metric_Provenance.pdf for row-level sources."
        ),
        (
            "Next. "
            "PeBRx workstream starting; dashboard refreshes when new harmonized exports and REDCap pulls are dropped in data/."
        ),
    ]

    for i, text in enumerate(bullets):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = text
        para.level = 0
        para.font.size = Pt(13)
        para.font.color.rgb = GRAY
        para.space_after = Pt(10)

    foot = slide.shapes.add_textbox(Inches(0.6), h - Inches(0.55), w - Inches(1.2), Inches(0.35))
    fp = foot.text_frame.paragraphs[0]
    fp.text = (
        f"Generated {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} · "
        "Data: James & Xinhui Wang (harmonized package); Adrienne REDCap exports Oct 2026"
    )
    fp.font.size = Pt(9)
    fp.font.color.rgb = GRAY
    fp.alignment = PP_ALIGN.LEFT


def _layout_named(prs: Presentation, *names: str):
    wanted = {n.lower() for n in names}
    for layout in prs.slide_layouts:
        if layout.name and layout.name.lower() in wanted:
            return layout
    return prs.slide_layouts[6]


def _delete_all_slides(prs: Presentation) -> None:
    for index in range(len(prs.slides) - 1, -1, -1):
        slide_id = prs.slides._sldIdLst[index]
        r_id = slide_id.rId
        prs.part.drop_rel(r_id)
        slide_id.getparent().remove(slide_id)


def _add_eac_title_slide(prs: Presentation, window: str, slide_count: int) -> None:
    layout = _layout_named(prs, "Title Slide", "1_Title Slide")
    slide = prs.slides.add_slide(layout)
    if slide.shapes.title:
        slide.shapes.title.text = "USC ADRC Clinical Core · Metrics Dashboard"
    stamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 1:
            ph.text = (
                f"Reporting window {window} · {slide_count} section captures · "
                f"{DECK_BUILD_ID} · {stamp}"
            )
            break


def _content_box(prs: Presentation) -> tuple[int, int, int, int]:
    """Full-width image area below a one-line caption (blank slide)."""
    left = EAC_MARGIN
    top = EAC_CAPTION_BAND
    width = prs.slide_width - EAC_MARGIN * 2
    height = prs.slide_height - EAC_CAPTION_BAND - EAC_BOTTOM
    return left, top, width, height


def _trim_near_white_borders(png: Path) -> tuple[int, int]:
    """Return (width, height) after trimming empty margins so centering reads true."""
    with Image.open(png) as im:
        rgb = im.convert("RGB")
        px = rgb.load()
        w, h = rgb.size
        threshold = 248

        def row_blank(y: int) -> bool:
            for x in range(w):
                r, g, b = px[x, y]
                if r < threshold or g < threshold or b < threshold:
                    return False
            return True

        def col_blank(x: int) -> bool:
            for y in range(h):
                r, g, b = px[x, y]
                if r < threshold or g < threshold or b < threshold:
                    return False
            return True

        top = 0
        while top < h and row_blank(top):
            top += 1
        bottom = h - 1
        while bottom > top and row_blank(bottom):
            bottom -= 1
        left = 0
        while left < w and col_blank(left):
            left += 1
        right = w - 1
        while right > left and col_blank(right):
            right -= 1

        if right <= left or bottom <= top:
            return w, h

        cropped = rgb.crop((left, top, right + 1, bottom + 1))
        cropped.save(png)
        return cropped.size


def _place_picture_fit(
    slide,
    png: Path,
    left: int,
    top: int,
    avail_w: int,
    avail_h: int,
) -> None:
    """Center image in slide content box (horizontal + vertical), scaled to fit."""
    iw, ih = _trim_near_white_borders(png)
    if iw <= 0 or ih <= 0:
        return
    scale = min(avail_w / iw, avail_h / ih)
    pic_w = int(iw * scale)
    pic_h = int(ih * scale)
    x = left + (avail_w - pic_w) // 2
    y = top + (avail_h - pic_h) // 2
    slide.shapes.add_picture(str(png), x, y, width=pic_w, height=pic_h)


def _add_eac_screenshot_slide(prs: Presentation, png: Path, caption: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    w = prs.slide_width
    cap = slide.shapes.add_textbox(EAC_MARGIN, Inches(0.1), w - EAC_MARGIN * 2, EAC_CAPTION_BAND)
    p = cap.text_frame.paragraphs[0]
    p.text = caption
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.alignment = PP_ALIGN.CENTER
    left, top, avail_w, avail_h = _content_box(prs)
    _place_picture_fit(slide, png, left, top, avail_w, avail_h)


def _add_section_divider_slide(
    prs: Presentation, title: str, subtitle: str = ""
) -> None:
    """Chapter slide so each dashboard section reads clearly in the deck."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    w, h = prs.slide_width, prs.slide_height
    box = slide.shapes.add_textbox(Inches(0.75), Inches(2.75), w - Inches(1.5), Inches(1.5))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.alignment = PP_ALIGN.CENTER
    if subtitle:
        sub = tf.add_paragraph()
        sub.text = subtitle
        sub.font.size = Pt(18)
        sub.font.color.rgb = PURPLE
        sub.alignment = PP_ALIGN.CENTER
        sub.space_before = Pt(8)


def _iter_screenshot_slides() -> list[dict]:
    return [row for row in load_screenshot_manifest() if (SCREENSHOT_DIR / row["file"]).is_file()]


def build_eac_template_deck(window: str) -> None:
    if not EAC_TEMPLATE.is_file():
        raise SystemExit(f"EAC template not found: {EAC_TEMPLATE}")

    shutil.copy2(EAC_TEMPLATE, OUTPUT_SCREENSHOTS_ONLY)
    prs = Presentation(str(OUTPUT_SCREENSHOTS_ONLY))
    _delete_all_slides(prs)
    slides = _iter_screenshot_slides()
    _add_eac_title_slide(prs, window, len(slides))
    prev_section: str | None = None
    for row in slides:
        sec = row.get("sectionId")
        if sec and sec != prev_section:
            _add_section_divider_slide(
                prs,
                str(row.get("sectionTitle") or sec),
                str(row.get("sectionSubtitle") or ""),
            )
            prev_section = sec
        _add_eac_screenshot_slide(prs, SCREENSHOT_DIR / row["file"], row["caption"])
    prs.save(str(OUTPUT_SCREENSHOTS_ONLY))
    stamped = (
        Path.home()
        / "Desktop"
        / f"USC_ADRC_Dashboard_Screenshots_{datetime.now().strftime('%Y-%m-%d_%H%M')}.pptx"
    )
    shutil.copy2(OUTPUT_SCREENSHOTS_ONLY, OUTPUT_SCREENSHOTS_DESKTOP)
    shutil.copy2(OUTPUT_SCREENSHOTS_ONLY, stamped)
    print(f"Also saved: {stamped}")


def add_image_slide(prs: Presentation, png: Path, caption: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    w, h = prs.slide_width, prs.slide_height
    margin = EAC_MARGIN
    box = slide.shapes.add_textbox(margin, Inches(0.25), w - margin * 2, Inches(0.4))
    p = box.text_frame.paragraphs[0]
    p.text = caption
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.alignment = PP_ALIGN.CENTER
    img_top = Inches(0.58)
    img_h = h - img_top - Inches(0.22) - EAC_BOTTOM
    _place_picture_fit(slide, png, margin, img_top, w - margin * 2, img_h)


def build_deck() -> None:
    m = load_metrics()
    capture_screenshots()

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_narrative_slide(prs, m)
    prev_section: str | None = None
    for row in _iter_screenshot_slides():
        sec = row.get("sectionId")
        if sec and sec != prev_section:
            _add_section_divider_slide(
                prs,
                str(row.get("sectionTitle") or sec),
                str(row.get("sectionSubtitle") or ""),
            )
            prev_section = sec
        add_image_slide(prs, SCREENSHOT_DIR / row["file"], row["caption"])

    OUTPUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT_PPTX))
    shutil.copy2(OUTPUT_PPTX, OUTPUT_DOWNLOADS)

    build_eac_template_deck(m["window"])

    print(f"EAC template deck: {OUTPUT_SCREENSHOTS_DESKTOP}")
    print(f"Narrative deck: {OUTPUT_DOWNLOADS}")
    print(f"Screenshots: {SCREENSHOT_DIR}")


def main() -> int:
    build_deck()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
